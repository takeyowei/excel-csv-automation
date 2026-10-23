#!/usr/bin/env python3
"""
merge_reports.py - Merge messy CSV/Excel exports into one clean Excel report.

Typical use case: a shop exports sales data from several systems (or months)
and the files never quite match - column names differ ("Order ID" vs
"order_id" vs "Order Id"), there are duplicate rows, stray spaces and blank
cells. This tool:

1. Reads any mix of .csv / .xlsx / .xls files (files or whole folders).
2. Normalises column names to snake_case and maps known aliases (qty -> quantity).
3. Cleans the data: trims whitespace, treats empty strings as missing,
   drops fully-empty rows and exact duplicates.
4. Writes a formatted Excel report with two sheets:
   - "Clean Data": the merged, cleaned rows (with their source file).
   - "Summary":    per-product totals (orders, quantity, revenue).

Usage
-----
    python merge_reports.py sample_data                 # demo data
    python merge_reports.py jan.csv feb.xlsx --output output/report.xlsx
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xls"}

# Column names that mean the same thing across different exports.
COLUMN_ALIASES = {
    "qty": "quantity",
    "id": "order_id",
    "price": "unit_price",
    "cost": "unit_price",
    "date": "order_date",
}

# Columns that should be numeric after cleaning (when present).
NUMERIC_COLUMNS = {"order_id", "quantity", "unit_price"}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY_FORMAT = "#,##0.00"


# ---------------------------------------------------------------------------
# Loading & cleaning
# ---------------------------------------------------------------------------

def normalize_column(name: str) -> str:
    """'Order ID ' / 'Unit-Price' / 'order_id' -> 'order_id' / 'unit_price'."""
    col = re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")
    col = re.sub(r"_+", "_", col)
    return COLUMN_ALIASES.get(col, col)


def load_file(path: Path) -> pd.DataFrame:
    """Read one CSV/Excel file and tag each row with its source file name."""
    if path.suffix.lower() == ".csv":
        df = pd.read_csv(path, dtype=str, keep_default_na=False)
    else:
        df = pd.read_excel(path, dtype=str)
        df = df.fillna("")
    df.columns = [normalize_column(c) for c in df.columns]
    df["source_file"] = path.name
    return df


def clean_data(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Clean a merged dataframe. Returns (clean_df, stats)."""
    stats = {"rows_in": len(df)}

    # Trim whitespace in every cell; empty strings count as missing.
    df = df.map(lambda v: v.strip() if isinstance(v, str) else v)
    df = df.replace("", pd.NA)

    # Drop rows that carry no data at all (ignoring the source tag).
    data_cols = [c for c in df.columns if c != "source_file"]
    df = df.dropna(how="all", subset=data_cols)

    # Drop exact duplicate records (same values, regardless of source file).
    before = len(df)
    df = df.drop_duplicates(subset=data_cols, keep="first")
    stats["duplicates_removed"] = before - len(df)

    # Numeric columns: convert, unparseable values become missing.
    for col in NUMERIC_COLUMNS & set(df.columns):
        df[col] = pd.to_numeric(df[col], errors="coerce")

    stats["rows_out"] = len(df)
    return df.reset_index(drop=True), stats


def build_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Per-product totals: orders, total quantity, total revenue."""
    required = {"product", "quantity", "unit_price"}
    if not required.issubset(df.columns):
        return pd.DataFrame()  # not sales-like data; summary not applicable
    work = df.copy()
    work["revenue"] = work["quantity"].fillna(0) * work["unit_price"].fillna(0)
    summary = (
        work.groupby("product", dropna=False)
        .agg(orders=("product", "size"),
             total_quantity=("quantity", "sum"),
             total_revenue=("revenue", "sum"))
        .reset_index()
        .sort_values("total_revenue", ascending=False)
    )
    total_row = {
        "product": "TOTAL",
        "orders": int(summary["orders"].sum()),
        "total_quantity": summary["total_quantity"].sum(),
        "total_revenue": summary["total_revenue"].sum(),
    }
    return pd.concat([summary, pd.DataFrame([total_row])], ignore_index=True)


# ---------------------------------------------------------------------------
# Excel output with formatting
# ---------------------------------------------------------------------------

def _style_sheet(ws, money_columns: set[str]) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    headers = [c.value for c in ws[1]]
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        values = [str(c.value) for c in ws[letter] if c.value is not None]
        ws.column_dimensions[letter].width = min(max(len(v) for v in values) + 2, 42)
        if header in money_columns:
            for cell in ws[letter][1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = MONEY_FORMAT


def write_report(clean: pd.DataFrame, summary: pd.DataFrame, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        clean.to_excel(writer, sheet_name="Clean Data", index=False)
        if not summary.empty:
            summary.to_excel(writer, sheet_name="Summary", index=False)
        money = {"unit_price", "total_revenue"}
        for sheet in writer.sheets.values():
            _style_sheet(sheet, money)


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def collect_files(paths: list[str]) -> list[Path]:
    files: list[Path] = []
    for raw in paths:
        path = Path(raw)
        if path.is_dir():
            files.extend(sorted(p for p in path.iterdir()
                                  if p.suffix.lower() in SUPPORTED_SUFFIXES))
        elif path.suffix.lower() in SUPPORTED_SUFFIXES:
            files.append(path)
        else:
            raise ValueError(f"unsupported input: {raw} (use .csv/.xlsx files or a folder)")
    if not files:
        raise ValueError("no input files found")
    return files


def merge_files(paths: list[str], output_path: Path) -> dict:
    files = collect_files(paths)
    frames = [load_file(f) for f in files]
    merged = pd.concat(frames, ignore_index=True)
    clean, stats = clean_data(merged)
    summary = build_summary(clean)
    write_report(clean, summary, output_path)
    return {"files": [f.name for f in files], "stats": stats,
            "summary_rows": len(summary), "output": output_path}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Merge and clean CSV/Excel exports "
                                                 "into a formatted Excel report.")
    parser.add_argument("inputs", nargs="+", help="files and/or folders to merge")
    parser.add_argument("--output", default="output/merged_report.xlsx",
                        help="report path (default: output/merged_report.xlsx)")
    args = parser.parse_args(argv)
    try:
        result = merge_files(args.inputs, Path(args.output))
    except (ValueError, OSError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    stats = result["stats"]
    print(f"Files merged : {', '.join(result['files'])}")
    print(f"Rows in      : {stats['rows_in']}")
    print(f"Duplicates   : {stats['duplicates_removed']} removed")
    print(f"Rows out     : {stats['rows_out']}")
    print(f"Report saved : {result['output']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
